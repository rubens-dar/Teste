from __future__ import annotations

import csv
import io
import re
from pathlib import Path
from urllib.parse import urlsplit, parse_qs

from leadhunter.config import CONFIG
from leadhunter.models import (
    ADDRESS, CNPJ, EMAIL, FANTASIA, INSTAGRAM, OWNER, PHONE, RAZAO, WEBSITE,
    WHATSAPP, Lead, Result,
)
from leadhunter.normalize import format_cnpj, pretty_phone, slug_name

NAME_KEYS = ("nome do vendedor", "nome vendedor", "vendedor", "nome da loja", "loja",
             "nickname", "apelido", "seller", "nome", "razao social", "empresa")
URL_KEYS = ("link do perfil", "link perfil", "perfil", "link", "url", "permalink",
            "pagina", "site do vendedor")
CITY_KEYS = ("cidade", "municipio", "city")
UF_KEYS = ("uf", "estado", "state")
ID_KEYS = ("seller_id", "seller id", "id do vendedor", "id vendedor", "cust_id", "id")


def _norm_header(h: str) -> str:
    txt = (h or "").strip().lower()
    txt = re.sub(r"[^\w\s]", " ", txt)
    return re.sub(r"\s+", " ", txt).strip()


def _pick(row: dict, keys: tuple[str, ...]) -> str:
    for key in keys:
        for header, value in row.items():
            if header == key and str(value or "").strip():
                return str(value).strip()
    for key in keys:
        for header, value in row.items():
            if key in header and str(value or "").strip():
                return str(value).strip()
    return ""


def _sheets_export_url(url: str) -> str | None:
    """Converte um link de Google Sheets no equivalente que exporta CSV."""
    parts = urlsplit(url)
    if "docs.google.com" not in parts.netloc or "/spreadsheets/" not in parts.path:
        return None
    match = re.search(r"/spreadsheets/d/([A-Za-z0-9_\-]+)", parts.path)
    if not match:
        return None
    doc_id = match.group(1)
    gid = "0"
    frag = parse_qs(parts.fragment or "")
    if frag.get("gid"):
        gid = frag["gid"][0]
    elif parse_qs(parts.query).get("gid"):
        gid = parse_qs(parts.query)["gid"][0]
    return f"https://docs.google.com/spreadsheets/d/{doc_id}/export?format=csv&gid={gid}"


def _rows_from_csv(text: str) -> list[dict]:
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        delim = dialect.delimiter
    except csv.Error:
        delim = ";" if sample.count(";") > sample.count(",") else ","
    reader = csv.DictReader(io.StringIO(text), delimiter=delim)
    return [{_norm_header(k): v for k, v in row.items() if k} for row in reader]


def _rows_from_xlsx(path: Path) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    headers = [_norm_header(str(h or "")) for h in next(rows, [])]
    out = []
    for values in rows:
        if not any(v is not None and str(v).strip() for v in values):
            continue
        out.append({h: ("" if v is None else str(v)) for h, v in zip(headers, values) if h})
    wb.close()
    return out


def load_leads(source: str) -> list[Lead]:
    """Aceita caminho de CSV/XLSX, URL de CSV ou link do Google Sheets."""
    if source.startswith("http"):
        from leadhunter.http import fetch

        url = _sheets_export_url(source) or source
        status, body = fetch(url, cache_ttl=0)
        if status != 200 or not body.strip():
            raise SystemExit(
                f"Nao consegui baixar a planilha ({status}). "
                "Se for Google Sheets, marque 'qualquer pessoa com o link pode ver' "
                "ou exporte para CSV e passe o caminho do arquivo."
            )
        rows = _rows_from_csv(body)
    else:
        path = Path(source).expanduser()
        if not path.exists():
            raise SystemExit(f"Arquivo nao encontrado: {path}")
        rows = _rows_from_xlsx(path) if path.suffix.lower() in (".xlsx", ".xlsm") \
            else _rows_from_csv(path.read_text(encoding="utf-8-sig", errors="replace"))

    leads: list[Lead] = []
    for i, row in enumerate(rows, start=2):
        lead = Lead(
            row=i,
            nome_ml=_pick(row, NAME_KEYS),
            perfil_url=_pick(row, URL_KEYS),
            cidade=_pick(row, CITY_KEYS) or CONFIG.cidade,
            uf=_pick(row, UF_KEYS) or CONFIG.uf,
            seller_id=re.sub(r"\D", "", _pick(row, ID_KEYS)),
            extra={k: v for k, v in row.items() if v},
        )
        if lead.nome_ml or lead.perfil_url or lead.seller_id:
            leads.append(lead)
    if not leads:
        raise SystemExit(
            "Nenhum lead lido. Cabecalhos encontrados: "
            + ", ".join(sorted({k for r in rows for k in r})) or "(vazio)"
        )
    return leads


COLUMNS = [
    "linha", "nome_ml", "perfil_ml", "seller_id",
    "cnpj", "razao_social", "nome_fantasia",
    "proprietario", "socios",
    "whatsapp_1", "whatsapp_2", "telefone_1", "telefone_2",
    "email_1", "email_2",
    "instagram", "site", "endereco",
    "tier", "score", "status", "fontes", "erros",
]


def _row(res: Result) -> dict:
    lead = res.lead
    whats = [pretty_phone(e.value) for e in res.of(WHATSAPP)]
    phones = [pretty_phone(e.value) for e in res.of(PHONE) if
              pretty_phone(e.value) not in whats]
    emails = [e.value for e in res.of(EMAIL)]
    owners = [e.value for e in res.of(OWNER)]
    cnpj = res.first(CNPJ)

    has_contact = bool(whats or phones or emails)
    status = "OK" if (owners and has_contact) else (
        "SO CONTATO EMPRESA" if has_contact else (
            "SO IDENTIFICACAO" if cnpj or owners else "SEM DADOS"))

    return {
        "linha": lead.row,
        "nome_ml": lead.nome_ml,
        "perfil_ml": lead.perfil_url,
        "seller_id": lead.seller_id,
        "cnpj": format_cnpj(cnpj) if cnpj else "",
        "razao_social": res.first(RAZAO),
        "nome_fantasia": res.first(FANTASIA),
        "proprietario": owners[0] if owners else "",
        "socios": " | ".join(owners),
        "whatsapp_1": whats[0] if whats else "",
        "whatsapp_2": whats[1] if len(whats) > 1 else "",
        "telefone_1": phones[0] if phones else "",
        "telefone_2": phones[1] if len(phones) > 1 else "",
        "email_1": emails[0] if emails else "",
        "email_2": emails[1] if len(emails) > 1 else "",
        "instagram": res.first(INSTAGRAM),
        "site": res.first(WEBSITE),
        "endereco": res.first(ADDRESS),
        "tier": res.tier,
        "score": round(res.score, 1),
        "status": status,
        "fontes": ", ".join(res.sources()),
        "erros": " | ".join(res.errors[:3]),
    }


def write_outputs(results: list[Result], outdir: Path) -> dict[str, Path]:
    outdir.mkdir(parents=True, exist_ok=True)
    rows = [_row(r) for r in results]

    csv_path = outdir / "leads_enriquecidos.csv"
    with csv_path.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=COLUMNS)
        writer.writeheader()
        writer.writerows(rows)

    manual = [r for r in rows if r["status"] in ("SEM DADOS", "SO IDENTIFICACAO")]
    manual_path = outdir / "revisao_manual.csv"
    with manual_path.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=COLUMNS)
        writer.writeheader()
        writer.writerows(manual)

    paths = {"csv": csv_path, "revisao": manual_path}

    try:
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "leads"
        ws.append(COLUMNS)
        for row in rows:
            ws.append([row[c] for c in COLUMNS])
        ws.freeze_panes = "A2"
        for idx, col in enumerate(COLUMNS, start=1):
            width = max(len(col), *(len(str(r[col])) for r in rows)) if rows else len(col)
            ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = min(width + 2, 45)
        xlsx_path = outdir / "leads_enriquecidos.xlsx"
        wb.save(xlsx_path)
        paths["xlsx"] = xlsx_path
    except Exception:  # openpyxl ausente nao deve derrubar a execucao
        pass

    import json

    json_path = outdir / "evidencias.json"
    json_path.write_text(
        json.dumps([r.to_dict() for r in results], ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    paths["json"] = json_path
    return paths


def dedupe_key(res: Result) -> str:
    return res.first(CNPJ) or slug_name(res.first(RAZAO) or res.lead.nome_ml)
